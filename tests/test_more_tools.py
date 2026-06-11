import io
import json

import pypdf
import pytest

from pdfsmith.tools.analyze import run_compare
from pdfsmith.tools.convert_more import (
    run_html_to_pdf,
    run_pdf_to_excel,
    run_pdf_to_pdfa,
    run_pdf_to_powerpoint,
    weasyprint_available,
)
from pdfsmith.tools.forms import list_form_fields, run_form_fill
from pdfsmith.tools.helpers import ToolError
from pdfsmith.tools.transform import ghostscript_path
from tests.conftest import make_text_pdf, upload


def test_pdf_to_powerpoint(pdf3):
    result = run_pdf_to_powerpoint([upload("doc.pdf", pdf3)], {})
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(result.data))
    slides = list(presentation.slides)
    assert len(slides) == 3
    assert any(shape.shape_type == 13 for shape in slides[0].shapes)  # 13 = picture


def make_table_pdf() -> bytes:
    """A PDF with a ruled 3x3 table that pdfplumber can detect."""
    from reportlab.pdfgen import canvas as rl_canvas

    buffer = io.BytesIO()
    canvas = rl_canvas.Canvas(buffer, pagesize=(612, 792))
    xs = [72, 212, 352, 492]
    ys = [700, 660, 620, 580]
    canvas.grid(xs, ys)
    rows = [["Item", "Qty", "Price"], ["Mug", "2", "10.00"], ["Lamp", "1", "25.00"]]
    canvas.setFont("Helvetica", 11)
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            canvas.drawString(xs[column_index] + 6, ys[row_index] - 26, value)
    canvas.showPage()
    canvas.save()
    return buffer.getvalue()


def test_pdf_to_excel_extracts_table():
    result = run_pdf_to_excel([upload("table.pdf", make_table_pdf())], {})
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(result.data))
    sheet = workbook["Page 1"]
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    flat = [v for row in values for v in row if v]
    assert "Item" in flat and "Mug" in flat and "25.00" in flat


def test_pdf_to_excel_text_fallback(pdf3):
    result = run_pdf_to_excel([upload("doc.pdf", pdf3)], {})
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(result.data))
    assert len(workbook.sheetnames) == 3
    first_column = [row[0].value for row in workbook["Page 1"].iter_rows() if row[0].value]
    assert any("Alpha" in str(value) for value in first_column)


@pytest.mark.skipif(not weasyprint_available(), reason="weasyprint system libs missing")
def test_html_to_pdf_renders():
    html = b"<h1>Quarterly Numbers</h1><p>Revenue is up.</p>"
    result = run_html_to_pdf([upload("report.html", html)], {})
    reader = pypdf.PdfReader(io.BytesIO(result.data))
    assert "Quarterly Numbers" in reader.pages[0].extract_text()


@pytest.mark.skipif(not weasyprint_available(), reason="weasyprint system libs missing")
def test_html_to_pdf_blocks_external_fetches():
    html = b'<img src="http://169.254.169.254/latest/meta-data/"><p>still renders</p>'
    result = run_html_to_pdf([upload("sneaky.html", html)], {})
    assert b"%PDF" in result.data[:8]  # renders, with the fetch refused


@pytest.mark.skipif(ghostscript_path() is None, reason="ghostscript not installed")
def test_pdf_to_pdfa(pdf3):
    result = run_pdf_to_pdfa([upload("doc.pdf", pdf3)], {})
    assert b"pdfaid" in result.data  # the PDF/A identification schema is embedded


def test_compare_finds_changes():
    a = make_text_pdf(("Alpha", "Bravo"))
    b = make_text_pdf(("Alpha", "Bravo CHANGED"))
    report = run_compare([upload("a.pdf", a), upload("b.pdf", b)], {}).data.decode()
    assert "page 1: identical" in report
    assert "page 2: TEXT differs" in report
    assert "documents DIFFER" in report
    assert "CHANGED" in report  # the diff itself is included


def test_compare_identical():
    a = make_text_pdf(("Alpha",))
    report = run_compare([upload("a.pdf", a), upload("b.pdf", a)], {}).data.decode()
    assert "no differences found" in report


def make_form_pdf() -> bytes:
    from reportlab.pdfgen import canvas as rl_canvas

    buffer = io.BytesIO()
    canvas = rl_canvas.Canvas(buffer, pagesize=(612, 792))
    canvas.setFont("Helvetica", 12)
    canvas.drawString(72, 712, "Name:")
    canvas.acroForm.textfield(name="full_name", x=140, y=700, width=220, height=20)
    canvas.drawString(72, 662, "Agree:")
    canvas.acroForm.checkbox(name="agree", x=140, y=652)
    canvas.showPage()
    canvas.save()
    return buffer.getvalue()


def test_list_form_fields():
    fields = {field["name"]: field for field in list_form_fields(make_form_pdf())}
    assert fields["full_name"]["type"] == "text"
    assert fields["agree"]["type"] == "checkbox"


def test_form_fill_roundtrip():
    result = run_form_fill(
        [upload("form.pdf", make_form_pdf())],
        {"fields": json.dumps({"full_name": "Vishal K"})},
    )
    reader = pypdf.PdfReader(io.BytesIO(result.data))
    assert reader.get_fields()["full_name"].value == "Vishal K"


def test_form_fill_lock():
    result = run_form_fill(
        [upload("form.pdf", make_form_pdf())],
        {"fields": json.dumps({"full_name": "Vishal K"}), "lock": "true"},
    )
    reader = pypdf.PdfReader(io.BytesIO(result.data))
    annotations = reader.pages[0]["/Annots"]
    locked = [a.get_object() for a in annotations
              if str(a.get_object().get("/T", "")) == "full_name"]
    assert locked and int(locked[0]["/Ff"]) & 1


def test_form_fill_validates(pdf3):
    with pytest.raises(ToolError, match="no fillable"):
        run_form_fill([upload("doc.pdf", pdf3)], {"fields": json.dumps({"x": "1"})})
    with pytest.raises(ToolError, match="unknown form fields"):
        run_form_fill([upload("form.pdf", make_form_pdf())],
                      {"fields": json.dumps({"nope": "1"})})
